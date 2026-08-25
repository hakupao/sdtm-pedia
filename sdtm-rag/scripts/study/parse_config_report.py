"""ConfigurationReport (Viedoc 导出 xlsx) 确定性解析器. 零 LLM."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def read_sheet_records(ws, *, has_section_row: bool = True) -> list[dict[str, Any]]:
    """表头 → [{'分组::列名': str, '_row': int}]; 全空行跳过, 短行补空串.

    has_section_row=True:  三行表头 (行1 sheet名 / 行2 分组, 空白前向填充 / 行3 列名), 数据自行 4 起.
    has_section_row=False: 两行表头 (行1 sheet名 / 行2 列名), 分组名取 sheet 名, 数据自行 3 起.
                           真实 Code lists sheet 即此形态.
    """
    rows = ws.iter_rows(values_only=True)
    try:
        next(rows)                               # 行1: sheet 名, 丢弃
        sections_raw: tuple = next(rows) if has_section_row else ()
        colnames = next(rows)
    except StopIteration:
        raise ValueError(
            f"sheet {ws.title!r}: header rows missing "
            f"(expected {3 if has_section_row else 2}-row header)"
        ) from None
    keys: list[str] = []
    cur = "" if has_section_row else ws.title
    for i, c in enumerate(colnames):
        section = _cell(sections_raw[i]) if i < len(sections_raw) else ""
        if section:
            cur = section
        keys.append(f"{cur}::{_cell(c)}")
    first_data_row = 4 if has_section_row else 3
    records: list[dict[str, Any]] = []
    for i, row in enumerate(rows, start=first_data_row):
        values = [_cell(v) for v in row[: len(keys)]]
        values += [""] * (len(keys) - len(values))   # 短行补齐, 否则 zip 静默丢 key
        if not any(values):
            continue
        rec: dict[str, Any] = dict(zip(keys, values))
        rec["_row"] = i
        records.append(rec)
    return records


@dataclass(frozen=True)
class FormDef:
    oid: str
    name: str
    summary_format: str
    description: str
    in_use: str
    row: int
    is_trailer: bool = False   # Viedoc 表尾脚注行 (oid 含空格 / In use 空), 真实 form 21 个均非此形态


@dataclass(frozen=True)
class ItemRow:
    row: int
    form_oid: str
    form_name: str
    row_type: str          # 'Item group' | 'Item' | 'Trailer' (表尾脚注, 白名单外归一化)
    group_oid: str
    group_name: str
    item_oid: str
    data_type: str
    required: bool
    min_length: str
    max_length: str
    data_checks: str
    system_checks: str
    label: str
    control_type: str
    choices: str
    unit: str
    description: str
    instructions: str
    visible_condition: str
    output_field_id: str
    output_field_label: str
    raw: dict = field(repr=False)


@dataclass(frozen=True)
class Codelist:
    oid: str
    data_type: str
    entries: list


def _req(rec: dict, key: str) -> str:
    if key not in rec:
        raise KeyError(f"{key} (available: {sorted(k for k in rec if k != '_row')[:8]}...)")
    return rec[key]


def parse_forms(path: Path) -> list[FormDef]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        records = read_sheet_records(wb["Forms"])
    finally:
        wb.close()      # read_only 模式持有文件句柄
    out: list[FormDef] = []
    for r in records:
        oid = _req(r, "General::Id")
        in_use = r.get("General::In use", "")
        out.append(FormDef(
            oid=oid, name=r.get("General::Name", ""),
            summary_format=r.get("General::Summary format", ""),
            description=r.get("General::Description", ""),
            in_use=in_use, row=r["_row"],
            is_trailer=(" " in oid or not in_use),
        ))
    return out


def parse_items(path: Path) -> list[ItemRow]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        records = read_sheet_records(wb["Items and Groups"])
    finally:
        wb.close()
    out: list[ItemRow] = []
    for r in records:
        raw_type = r.get("Type and container::Field type", "")
        out.append(ItemRow(
            row=r["_row"],
            form_oid=_req(r, "Type and container::Form ID"),
            form_name=r.get("Type and container::Form Name", ""),
            row_type=raw_type if raw_type in ("Item", "Item group") else "Trailer",
            group_oid=r.get("Type and container::Item group ID", ""),
            group_name=r.get("Type and container::Item group name", ""),
            item_oid=_req(r, "Validation::Item ID"),
            data_type=r.get("Validation::Data type", ""),
            required=r.get("Validation::Required field", "") == "X",
            min_length=r.get("Validation::Minimum length", ""),
            max_length=r.get("Validation::Max length", ""),
            data_checks=r.get("Validation::Data checks", ""),
            system_checks=r.get("Validation::System checks", ""),
            label=r.get("General::Field label", ""),
            control_type=r.get("General::Control Type", ""),
            choices=r.get("General::Choices", ""),
            unit=r.get("General::Measurement Unit", ""),
            description=r.get("General::Description", ""),
            instructions=r.get("General::Instructions for user", ""),
            visible_condition=r.get("Visibility::Show on simple condition", ""),
            output_field_id=r.get("Output::Output Field ID", ""),
            output_field_label=r.get("Output::Output Field Label", ""),
            raw={k: v for k, v in r.items() if k != "_row"},
        ))
    return out


def parse_codelists(path: Path) -> dict[str, Codelist]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        records = read_sheet_records(wb["Code lists"], has_section_row=False)
    finally:
        wb.close()
    grouped: dict[str, Codelist] = {}
    for r in records:
        oid = _req(r, "Code lists::OID")
        cl = grouped.setdefault(
            oid, Codelist(oid=oid, data_type=r.get("Code lists::Data Type", ""), entries=[])
        )
        cl.entries.append((r.get("Code lists::Code value", ""), r.get("Code lists::Code text", "")))
    return grouped


# ── Study workflow 三表 ────────────────────────────────────────────────
# 三个 sheet 在同一个 ConfigReport 里, 但表头形态不同 (实测):
#   Events / Activities  = 三行表头 (has_section_row=True),  数据自行 4 起
#   Forms                = 两行表头 (has_section_row=False), 数据自行 3 起
# 三表末尾均有 Viedoc 说明性脚注行, 形态: ID 列为空或含空格 (真 OID 无空格)。
# 与 FormDef.is_trailer 同惯例: 解析器**返回全部行**, 过滤交给调用方 (台账要记脚注)。


@dataclass(frozen=True)
class EventDef:
    oid: str
    name: str
    description: str
    event_type: str
    visibility_condition: str
    sched_reference: str
    sched_minus_days: str
    sched_plus_days: str
    row: int
    is_trailer: bool = False


@dataclass(frozen=True)
class ActivityDef:
    oid: str
    event_oid: str
    event_name: str
    name: str
    description: str
    visibility_condition: str
    row: int
    is_trailer: bool = False


@dataclass(frozen=True)
class FormAssignment:
    event_oid: str
    event_name: str
    activity_oid: str
    activity_name: str
    form_oid: str
    repeating: str
    item_visibility: str
    hidden_items: str
    row: int
    is_trailer: bool = False


def _read(path: Path, sheet: str, *, has_section_row: bool) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return read_sheet_records(wb[sheet], has_section_row=has_section_row)
    finally:
        wb.close()      # read_only 模式持有文件句柄


def _is_footnote(oid: str) -> bool:
    """脚注判据: ID 列为空或含空格。真 OID 是无空格标识符, 脚注是整句说明文字。"""
    return (not oid) or (" " in oid)


def parse_events(path: Path) -> list[EventDef]:
    out: list[EventDef] = []
    for r in _read(path, "Study workflow-Events", has_section_row=True):
        oid = _req(r, "General::Study event ID")
        out.append(EventDef(
            oid=oid,
            name=r.get("General::Event name", ""),
            description=r.get("General::Study event description", ""),
            event_type=r.get("General::Event type", ""),
            visibility_condition=r.get("Visibility::Visibility condition", ""),
            sched_reference=r.get("Scheduling::Reference", ""),
            sched_minus_days=r.get("Scheduling::- days", ""),
            sched_plus_days=r.get("Scheduling::+ days", ""),
            row=r["_row"],
            is_trailer=_is_footnote(oid),
        ))
    return out


def parse_activities(path: Path) -> list[ActivityDef]:
    out: list[ActivityDef] = []
    for r in _read(path, "Study workflow-Activities", has_section_row=True):
        oid = _req(r, "General::Activity ID")
        out.append(ActivityDef(
            oid=oid,
            event_oid=r.get("General::Study event ID", ""),
            event_name=r.get("General::Event name", ""),
            name=r.get("General::Activity name", ""),
            description=r.get("General::Activity description", ""),
            visibility_condition=r.get("General::Visibility condition", ""),
            row=r["_row"],
            is_trailer=_is_footnote(oid),
        ))
    return out


def parse_form_assignments(path: Path) -> list[FormAssignment]:
    out: list[FormAssignment] = []
    for r in _read(path, "Study workflow-Forms", has_section_row=False):
        form_oid = _req(r, "Study workflow-Forms::Form ID")
        out.append(FormAssignment(
            event_oid=r.get("Study workflow-Forms::Event ID", ""),
            event_name=r.get("Study workflow-Forms::Event name", ""),
            activity_oid=r.get("Study workflow-Forms::Activity ID", ""),
            activity_name=r.get("Study workflow-Forms::Activity name", ""),
            form_oid=form_oid,
            repeating=r.get("Study workflow-Forms::Repeating", ""),
            item_visibility=r.get("Study workflow-Forms::Item visibility", ""),
            hidden_items=r.get("Study workflow-Forms::Hidden items", ""),
            row=r["_row"],
            # Forms 表脚注行的 Form ID 为空 (Event ID 列反而是整句说明文字)
            is_trailer=_is_footnote(form_oid),
        ))
    return out
