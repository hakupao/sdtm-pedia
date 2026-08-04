"""DEMO 导出 xlsx → item 实例值采样 (去重保序, 空值跳过). 零 LLM, 零模糊匹配.

真实 DEMO 结构 (v3 实测):
  sheet 名 = form OID (21/21 吻合); 每张 form sheet **两行表头** —— 行 1 = 日文 label 表头,
  行 2 = **Item OID 表头** (系统列位置为 SiteSeq/SubjectId/EventId 等列 ID), 数据自行 3 起。
  当前真实 DEMO 行 3 起为空 = **零数据行**。

v2 曾把行 2 当作唯一数据行采样, 导致 671 个"示例值"全部是该列自身的 OID 回声。v3 把行 2
识别为权威 OID 表头层, 并加反回声闸 (采样值 == 自身 OID 即判解析失败) 防止此类错误复发。

列解析分层 (均限定在 sheet 名给出的 form 作用域内):
  (a') 行 2 OID —— 权威层, 命中即定案 (比 label 字典硬);
  (a) 行 1 列头 == item_oid 直配 / (b) DEMO Items 字典 Label→ID / (c) form 域内 label 精确
      匹配 —— 行 2 缺位时的兜底; 行 2 在位时降为**独立佐证通道** (corroborated/conflict 计数)。
"""
from __future__ import annotations

import collections
import itertools
import sys
from pathlib import Path
from typing import Any

import openpyxl

_SKIP_SHEETS = {"README", "Items", "CodeLists"}
_ITEMS_SHEET = "Items"

# 行2 判为 OID 表头的确定性判据: 非空单元格 ∈ OID 全集 (catalog item_oid ∪ DEMO Items 字典
# ID) 的比例 ≥ 该阈值。真实 DEMO 实测 2387/2387 = 100%; 真实数据行 (日期/数值/日文自由文本)
# 不可能有 9 成的值恰好落在 OID 全集里, 故该阈值可确定性区分"OID 表头"与"数据行"两种形态。
_OID_HEADER_MIN_RATIO = 0.9


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _load_items_dict(wb) -> tuple[dict[str, set[str]], set[str]]:
    """DEMO 自带的 Items 字典 sheet (单行表头 ID/Label/...) → (label → {ID}, 全部 ID)."""
    if _ITEMS_SHEET not in wb.sheetnames:
        return {}, set()
    rows = wb[_ITEMS_SHEET].iter_rows(values_only=True)
    header = [_cell(c) for c in next(rows, [])]
    try:
        i_id, i_label = header.index("ID"), header.index("Label")
    except ValueError:
        return {}, set()
    by_label: dict[str, set[str]] = collections.defaultdict(set)
    all_ids: set[str] = set()
    for row in rows:
        oid = _cell(row[i_id]) if i_id < len(row) else ""
        label = _cell(row[i_label]) if i_label < len(row) else ""
        if oid:
            all_ids.add(oid)
        if oid and label:
            by_label[label].add(oid)
    return by_label, all_ids


def _is_oid_header(cells: list[str], universe: set[str]) -> bool:
    """行 2 是否为 OID 表头行 (判据见 _OID_HEADER_MIN_RATIO)."""
    non_empty = [c for c in cells if c]
    if not non_empty:
        return False
    return sum(1 for c in non_empty if c in universe) / len(non_empty) >= _OID_HEADER_MIN_RATIO


def sample_demo_values(
    demo_path: Path, catalog_items: list[dict], max_per_item: int = 5,
) -> tuple[dict[str, list[str]], dict]:
    """DEMO 各 form sheet 采样实例值.

    catalog_items: catalog.json 的 items (用 item_oid / form_oid / label 三键).
    返回 (item_oid → 去重实例值列表, resolution stats). 分层与判据见模块 docstring.
    兜底路给出恰 1 个候选即解析; >1 记 ambiguous 并跳过该列; 全不中记 unmatched.
    采样值 == 该列自身 OID 时判为解析失败, 剔除并计入 echo_dropped.
    """
    scope_by_form: dict[str, set[str]] = collections.defaultdict(set)
    label_by_form: dict[str, dict[str, set[str]]] = collections.defaultdict(
        lambda: collections.defaultdict(set)
    )
    all_oids: set[str] = set()
    for it in catalog_items:
        oid, form = _cell(it.get("item_oid")), _cell(it.get("form_oid"))
        if not oid or not form:
            continue
        all_oids.add(oid)
        scope_by_form[form].add(oid)
        label = _cell(it.get("label"))
        if label:
            label_by_form[form][label].add(oid)

    samples: dict[str, list[str]] = {}
    stats: dict[str, Any] = {
        "resolved_row2_oid": 0, "resolved_direct": 0, "resolved_items_dict": 0,
        "resolved_label": 0, "corroborated": 0, "conflict_row2_label": 0,
        "ambiguous": 0, "unmatched": 0, "unknown_sheets": 0, "oid_header_sheets": 0,
        "data_rows": 0, "echo_dropped": 0, "echo_dropped_degraded": 0,
        "per_sheet": {}, "per_sheet_data_rows": {},
    }

    wb = openpyxl.load_workbook(demo_path, read_only=True)
    try:
        items_dict, dict_ids = _load_items_dict(wb)
        universe = all_oids | dict_ids
        for ws in wb.worksheets:
            if ws.title in _SKIP_SHEETS:
                continue
            if ws.title not in scope_by_form:
                stats["unknown_sheets"] += 1   # sheet 名不是 catalog form OID: 静默失效哨
            scope = scope_by_form.get(ws.title, set())
            rows = ws.iter_rows(values_only=True)
            label_header = [_cell(c) for c in next(rows, ())]

            # 行 2 可能是 OID 表头, 也可能已是数据 —— 判定后决定数据自哪行起
            peek_raw = next(rows, None)
            peek = [_cell(c) for c in peek_raw] if peek_raw is not None else []
            if peek_raw is not None and _is_oid_header(peek, universe):
                oid_header, data_iter = peek, rows
                stats["oid_header_sheets"] += 1
            elif peek_raw is not None:
                oid_header, data_iter = [], itertools.chain((peek_raw,), rows)
            else:
                oid_header, data_iter = [], rows

            col_oids: list[str | None] = []
            resolved_here = 0
            for j, name in enumerate(label_header):
                authoritative = ""
                if j < len(oid_header) and oid_header[j] in scope:
                    authoritative = oid_header[j]

                # 兜底/佐证三路 (先命中先定)
                fallback, fallback_key, fallback_ambiguous = None, "", False
                if name:
                    for key, cand in (
                        ("resolved_direct", {name} & scope),
                        ("resolved_items_dict", items_dict.get(name, set()) & scope),
                        ("resolved_label", label_by_form.get(ws.title, {}).get(name, set())),
                    ):
                        if not cand:
                            continue
                        if len(cand) == 1:
                            fallback, fallback_key = next(iter(cand)), key
                        else:
                            fallback_ambiguous = True
                        break

                if authoritative:
                    oid = authoritative
                    stats["resolved_row2_oid"] += 1
                    resolved_here += 1
                    if fallback:
                        stats["corroborated" if fallback == authoritative
                              else "conflict_row2_label"] += 1
                elif fallback:
                    oid = fallback
                    stats[fallback_key] += 1
                    resolved_here += 1
                else:
                    oid = None
                    if name:
                        stats["ambiguous" if fallback_ambiguous else "unmatched"] += 1
                col_oids.append(oid)
            stats["per_sheet"][ws.title] = resolved_here

            n_data = 0
            for row in data_iter:
                n_data += 1
                for oid, v in zip(col_oids, row):
                    if oid is None:
                        continue
                    sv = _cell(v)
                    if not sv:
                        continue
                    if sv == oid:
                        # 值 == 自身 OID: 行2 表头已检出的 sheet 里出现 = 疑似误杀真实数据
                        # (echo_dropped, 告警); 未检出的退化 sheet 里 = 闸拦下漏网表头行,
                        # 设计目标场景 (echo_dropped_degraded, 不告警)
                        stats["echo_dropped" if oid_header
                              else "echo_dropped_degraded"] += 1
                        continue
                    bucket = samples.setdefault(oid, [])
                    if sv not in bucket and len(bucket) < max_per_item:
                        bucket.append(sv)
            stats["data_rows"] += n_data
            stats["per_sheet_data_rows"][ws.title] = n_data
    finally:
        wb.close()      # read_only 模式持有文件句柄, build_field_cards 会循环调用
    if stats["echo_dropped"] and stats["data_rows"]:
        # 只对"表头已检出仍有值==OID"的形态告警 (真实数据被误杀的风险); 退化路径的
        # echo 是闸正确工作, 混进来会稀释信噪比 — 闸不放行, 但必须可见, 由人抽查裁决
        print(
            f"warning: anti-echo gate dropped {stats['echo_dropped']} value(s) "
            f"amid {stats['data_rows']} data rows — 可能误杀真实数据, 请抽查",
            file=sys.stderr,
        )
    return samples, stats
