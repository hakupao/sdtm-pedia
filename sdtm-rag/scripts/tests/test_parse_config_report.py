from pathlib import Path

import openpyxl
import pytest

from scripts.tests.study_fixtures import build_config_report
from scripts.study.parse_config_report import (
    parse_codelists,
    parse_forms,
    parse_items,
    read_sheet_records,
)


@pytest.fixture()
def report(tmp_path) -> Path:
    return build_config_report(tmp_path / "fake_report.xlsx")


def test_read_sheet_records_keys_and_forward_fill(report):
    wb = openpyxl.load_workbook(report, read_only=True)
    recs = read_sheet_records(wb["Items and Groups"])
    assert len(recs) == 4
    first = recs[0]
    assert first["Type and container::Form ID"] == "FAKEFORM1"
    # 分组前向填充: 第 2 列同属 Type and container
    assert first["Type and container::Form Name"] == "偽フォーム一"
    assert recs[1]["Validation::Item ID"] == "FAKEIT1"
    assert recs[1]["General::Choices"] == "CL_FAKE1"


def test_read_sheet_records_row_numbers(report):
    wb = openpyxl.load_workbook(report, read_only=True)
    recs = read_sheet_records(wb["Items and Groups"])
    assert recs[0]["_row"] == 4          # 数据从 xlsx 行 4 起
    assert recs[3]["_row"] == 7


def test_read_sheet_records_skips_blank_rows(tmp_path):
    p = build_config_report(tmp_path / "r.xlsx",
                            items_rows=[("", "") + ("",) * 19])
    wb = openpyxl.load_workbook(p, read_only=True)
    assert read_sheet_records(wb["Items and Groups"]) == []


def test_read_sheet_records_two_row_header(report):
    """真实 Code lists 是两行表头 (行1 sheet名 / 行2 列名), 分组名取 sheet 名."""
    wb = openpyxl.load_workbook(report, read_only=True)
    recs = read_sheet_records(wb["Code lists"], has_section_row=False)
    assert len(recs) == 3                     # 三行不能被表头吃掉一条
    assert recs[0]["_row"] == 3               # 数据从 xlsx 行 3 起
    assert recs[0]["Code lists::OID"] == "CL_FAKE1"
    assert recs[0]["Code lists::Code text"] == "偽選択肢はい"
    assert recs[2]["Code lists::OID"] == "CL_UNUSED"


class _StubSheet:
    """短行 sheet: openpyxl read_only 会自动补到 max_column, 故用 stub 直喂短 tuple."""

    title = "Stub"

    def __init__(self, rows):
        self._rows = rows

    def iter_rows(self, values_only=True):
        return iter(self._rows)


def test_read_sheet_records_pads_short_rows():
    ws = _StubSheet([
        ("Stub", "Stub", "Stub"),
        ("A", "A", "B"),
        ("c1", "c2", "c3"),
        ("v1",),                      # 短行: 缺 c2 / c3
    ])
    recs = read_sheet_records(ws)
    assert len(recs) == 1
    assert recs[0]["A::c1"] == "v1"
    assert recs[0]["A::c2"] == ""     # 缺列补空串, 不得静默丢 key
    assert recs[0]["B::c3"] == ""


def test_parse_forms(report):
    forms = parse_forms(report)
    assert [f.oid for f in forms] == ["FAKEFORM1", "FAKEFORM2"]
    assert forms[0].name == "偽フォーム一"


def test_parse_items_types_and_fields(report):
    items = parse_items(report)
    assert [r.row_type for r in items] == ["Item group", "Item", "Item", "Item"]
    it1 = items[1]
    assert (it1.form_oid, it1.item_oid, it1.data_type) == ("FAKEFORM1", "FAKEIT1", "integer")
    assert it1.required is True
    assert it1.choices == "CL_FAKE1"
    assert it1.label == "偽項目ラベル一"
    assert items[2].required is False
    assert items[2].visible_condition == "COND1"
    # group 行继承上下文: group_name 在后续 Item 行为空, 保留原值即可
    assert items[0].group_name == "グループ甲"
    # raw 保留全部列 (不丢信息)
    assert it1.raw["Output::Output Field ID"] == "OUT1"


def test_parse_codelists_grouping(report):
    cls = parse_codelists(report)
    assert set(cls) == {"CL_FAKE1", "CL_UNUSED"}
    assert cls["CL_FAKE1"].entries == [("1", "偽選択肢はい"), ("0", "偽選択肢いいえ")]
    assert cls["CL_FAKE1"].data_type == "integer"


def test_trailer_rows_flagged(tmp_path):
    """Viedoc 表尾脚注行: Forms 标 is_trailer, Items 归一化为 Trailer."""
    from scripts.tests.study_fixtures import DEFAULT_FORMS, DEFAULT_ITEMS
    trailer_form = ("See the Data checks sheet for details.", "", "", "", "")
    trailer_item = ("See the Data checks sheet for details.", "", "Footnote text", "", "") + ("",) * 16
    p = build_config_report(tmp_path / "t.xlsx",
                            forms_rows=list(DEFAULT_FORMS) + [trailer_form],
                            items_rows=list(DEFAULT_ITEMS) + [trailer_item])
    forms = parse_forms(p)
    assert [f.is_trailer for f in forms] == [False, False, True]
    items = parse_items(p)
    assert items[-1].row_type == "Trailer"
    # 非空非白名单值: 锁死"白名单归一化"语义, 防退化成 `raw_type or "Trailer"`
    # 含空格 = 脚注形态 (build_catalog 对无空格的未知值会 raise, 见 Task 5)
    assert items[-1].raw["Type and container::Field type"] == "Footnote text"


def test_parse_items_missing_column_raises(tmp_path):
    import openpyxl as _o
    p = tmp_path / "bad.xlsx"
    wb = _o.Workbook(); ws = wb.active; ws.title = "Items and Groups"
    ws.append(["Items and Groups"] * 2)
    ws.append(["Type and container", ""])
    ws.append(["Form ID", "Form Name"])
    ws.append(["F1", "n1"])
    wb.save(p)
    with pytest.raises(KeyError, match="Validation::Item ID"):
        parse_items(p)


# ---- P3: 短 sheet / 组合场景 / 句柄关闭 ----

def test_short_sheet_raises_value_error():
    """表头不足的 sheet 必须 ValueError 响亮失败, 而非裸 StopIteration."""
    from scripts.study.parse_config_report import read_sheet_records
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forms"
    ws.append(["Forms"])                     # 只有行1, 三行表头缺行2/行3
    with pytest.raises(ValueError, match="header"):
        read_sheet_records(ws)


def test_short_sheet_two_row_header_raises():
    from scripts.study.parse_config_report import read_sheet_records
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Code lists"
    ws.append(["Code lists"])                # 两行表头缺行2
    with pytest.raises(ValueError, match="header"):
        read_sheet_records(ws, has_section_row=False)


def test_two_row_header_short_row_padded():
    """组合场景: 两行表头 + 短数据行, 短行补空串不丢 key."""
    from scripts.study.parse_config_report import read_sheet_records
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Code lists"
    ws.append(["Code lists"] * 3)
    ws.append(["OID", "Code value", "Code text"])
    ws.append(["CL_FAKE1", "1"])             # 短行: 缺 Code text
    recs = read_sheet_records(ws, has_section_row=False)
    assert len(recs) == 1
    assert recs[0]["Code lists::Code text"] == ""


def test_parse_functions_close_workbook(tmp_path, monkeypatch):
    """read_only 模式持有文件句柄, parse_* 返回前必须关闭."""
    import scripts.study.parse_config_report as pcr
    path = build_config_report(tmp_path / "close.xlsx")
    closed: list[bool] = []
    real_load = pcr.openpyxl.load_workbook

    def spy_load(*a, **k):
        wb = real_load(*a, **k)
        real_close = wb.close
        def spy_close():
            closed.append(True)
            real_close()
        monkeypatch.setattr(wb, "close", spy_close, raising=False)
        return wb

    monkeypatch.setattr(pcr.openpyxl, "load_workbook", spy_load)
    pcr.parse_forms(path)
    pcr.parse_items(path)
    pcr.parse_codelists(path)
    assert len(closed) == 3
